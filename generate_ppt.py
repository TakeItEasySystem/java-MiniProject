import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY = RGBColor(15, 23, 42)       # #0F172A Dark Slate
    C_CARD_DARK = RGBColor(30, 41, 59)  # #1E293B
    C_BLUE = RGBColor(37, 99, 235)      # #2563EB Primary Blue
    C_BLUE_LIGHT = RGBColor(239, 246, 255) # Light blue card bg
    C_TEAL = RGBColor(13, 148, 136)     # #0D9488 Accent Teal
    C_EMERALD = RGBColor(16, 185, 129)  # #10B981 Green
    C_AMBER = RGBColor(245, 158, 11)    # #F59E0B Warning Amber
    C_BG_LIGHT = RGBColor(248, 250, 252)# #F8FAFC
    C_WHITE = RGBColor(255, 255, 255)
    C_BORDER = RGBColor(226, 232, 240)  # #E2E8F0
    C_TEXT_DARK = RGBColor(15, 23, 42)
    C_TEXT_MUTED = RGBColor(100, 116, 139)
    C_BORDER_DARK = RGBColor(51, 65, 85)

    def add_bg(slide, dark=False):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_NAVY if dark else C_BG_LIGHT
        bg.line.fill.background()
        return bg

    def add_header(slide, eyebrow, title, dark=False):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p1 = tf.paragraphs[0]
        p1.text = eyebrow.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = C_BLUE if not dark else C_TEAL
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_DARK if not dark else C_WHITE

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, dark=True)

    # Accent decorative top bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(1.5), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_BLUE
    bar.line.fill.background()

    # Title Box
    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "FULL STACK JAVA PROGRAMMING (FSJP) MINI PROJECT"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEAL
    p.space_after = Pt(12)

    p = tf1.add_paragraph()
    p.text = "WebEase: Proactive Web Presence & Client Operations Platform"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "An Enterprise-Grade, Multithreaded 'Website-as-a-Service' & Autonomous Telemetry System for MSMEs"
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(203, 213, 225)

    # Team Box (3 Cards)
    members = [
        ("Ayush Jha", "Roll No: A18", "System Architecture & Security"),
        ("Rishikant Mishra", "Roll No: A29", "Core Backend, Multithreading & DB"),
        ("Yadnesh Sawant", "Roll No: A68", "Frontend MVC & Telemetry UI")
    ]
    card_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    card_y = Inches(4.7)
    card_h = Inches(1.8)

    for i, (name, roll, role) in enumerate(members):
        x = start_x + i * (card_w + gap)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_DARK
        card.line.color.rgb = C_BORDER_DARK
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.2)
        tf.margin_left = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = name
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

        p2 = tf.add_paragraph()
        p2.text = roll
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_BLUE

        p3 = tf.add_paragraph()
        p3.text = role
        p3.font.size = Pt(11)
        p3.font.color.rgb = RGBColor(148, 163, 184)

    # Footer institution text
    ftr = s1.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.5))
    p = ftr.text_frame.paragraphs[0]
    p.text = "Thakur Shree DPS College  •  Department of Computer Science & Information Technology  •  Academic Year 2026"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(100, 116, 139)

    # =========================================================================
    # SLIDE 2: Problem Statement & Motivation
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Industry Context & Motivation", "The MSME Digital Dilemma & Why WebEase Exists")

    cards_data = [
        ("The Current Problem", C_NAVY, [
            "Over 63 Million Indian MSMEs lack a reliable digital footprint.",
            "Traditional agencies charge exorbitant upfront fees (₹50k+) with zero ongoing maintenance.",
            "Websites crash unnoticed, forms break, and SSL certificates expire without warning.",
            "Local owners (gyms, cafes, clinics) have zero technical bandwidth to manage hosting."
        ]),
        ("The Market Void", C_BLUE, [
            "DIY builders (Wix/WordPress) overwhelm non-technical owners with complex drag-and-drop tools.",
            "Agencies treat small businesses as one-off projects rather than ongoing managed services.",
            "No affordable solution provides proactive monitoring + instant WhatsApp support.",
            "Security vulnerability: Small websites rarely patch plugins or back up databases."
        ]),
        ("The WebEase Solution", C_TEAL, [
            "Website-as-a-Service (WaaS): Low setup + predictable monthly subscription.",
            "24/7 Autonomous Health Sweeps: Background daemon catches downtime before clients notice.",
            "Zero Password Vulnerability: Direct Google OAuth 2.0 via Firebase.",
            "Post-Order Omnichannel Onboarding: Instant connection via WhatsApp & Priority Calls."
        ])
    ]

    cw = Inches(3.64)
    for i, (title, color_top, bullets) in enumerate(cards_data):
        cx = Inches(0.8) + i * (cw + Inches(0.4))
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.8), cw, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        # Top indicator bar
        ib = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, Inches(1.8), cw, Inches(0.12))
        ib.fill.solid()
        ib.fill.fore_color.rgb = color_top
        ib.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.35)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(14)

        for b in bullets:
            p = tf.add_paragraph()
            p.text = "• " + b
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_MUTED
            p.space_after = Pt(10)

    # =========================================================================
    # SLIDE 3: Key Objectives & Features
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Functional Architecture", "Core System Objectives & Product Capabilities")

    feats = [
        ("24/7 Proactive Telemetry", "Autonomous Multithreaded Engine", "Continuous background daemon performs latency pings, HTTP diagnostics, and uptime scoring. Dispatches automated incident tickets on anomalies.", C_BLUE),
        ("Dynamic Polymorphic Pricing", "Object-Oriented Calculation Engine", "Polymorphic service tier generator with dynamic calculation for Starter, Growth, and Enterprise plans with real-time add-on configuration.", C_TEAL),
        ("Frictionless Google Auth", "Firebase OAuth 2.0 Integration", "Minimalist, zero-password authentication powered by Google Identity Services and Firebase JWT token validation on the backend servlet.", C_EMERALD),
        ("Operations Command Center", "Unified Administrative Console", "Live fleet dashboard displaying client website metrics, inbound audit leads, active orders, and maintenance ticketing system.", C_AMBER)
    ]

    fw = Inches(5.66)
    fh = Inches(2.35)
    for i, (title, sub, desc, color) in enumerate(feats):
        rx = Inches(0.8) + (i % 2) * (fw + Inches(0.4))
        ry = Inches(1.8) + (i // 2) * (fh + Inches(0.35))

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, ry, fw, fh)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.25)
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(12)
        p3.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 4: Technical Architecture & 3-Tier Model
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "System Design", "Three-Tier Enterprise Java Architecture")

    tiers = [
        ("1. Presentation Tier", "Client-Side & Views", C_BLUE, [
            "JSP 2.3 & JSTL 1.2 View Templates",
            "Responsive Vanilla CSS3 Design System",
            "AJAX Asynchronous Fetch API Requests",
            "Firebase JS SDK v10 (Google OAuth Popup)",
            "Client Consultation & Order Confirmation UI"
        ]),
        ("2. Application Tier", "Controllers & Business Logic", C_TEAL, [
            "Java Servlets (HttpServlet Controller Layer)",
            "AdminAuthFilter Session Security Filter",
            "MonitorService (Multithreaded Daemon)",
            "PricingService (Polymorphic Factory Pattern)",
            "Jakarta Mail & Brevo SMS Notification Service"
        ]),
        ("3. Data & Cloud Tier", "Persistence & Deployment", C_NAVY, [
            "SQLite 3 Embedded Relational Engine",
            "sqlite-jdbc-3.45.1.0 Driver in WAL Mode",
            "Thread-Safe Connection Management",
            "Apache Tomcat 9.0 Application Server",
            "Dockerized Cloud Hosting on Render.com"
        ])
    ]

    for i, (title, sub, col, items) in enumerate(tiers):
        tx = Inches(0.8) + i * (cw + Inches(0.4))
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, Inches(1.8), cw, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        ib = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, tx, Inches(1.8), cw, Inches(0.12))
        ib.fill.solid()
        ib.fill.fore_color.rgb = col
        ib.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.35)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.space_after = Pt(14)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_MUTED
            p.space_after = Pt(10)

    # =========================================================================
    # SLIDE 5: FSJP Syllabus Alignment (Units I to V)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Curriculum Mapping", "Alignment with Full Stack Java Programming (FSJP) Syllabus")

    units = [
        ("Unit I: Core Java & OOP", "ServicePlan (Abstract Base), StarterPlan, GrowthPlan, EnterprisePlan classes overriding calculation routines. Custom Exception hierarchy.", C_BLUE),
        ("Unit II: Multithreading & Net", "MonitorService implements Runnable daemon thread executing 45s loops. HttpURLConnection for network latency & response codes.", C_TEAL),
        ("Unit III: JDBC & Persistence", "SQLite JDBC driver loading, try-with-resources connection management, PRAGMA journal_mode=WAL, and DAO patterns.", C_EMERALD),
        ("Unit IV: Servlets & Filters", "CustomerAuthServlet, OrderServlet, SiteServlet, PricingServlet. AdminAuthFilter guarding protected management routes.", C_AMBER),
        ("Unit V: JSP, JSTL & Cloud", "Dynamic scriptless JSP views with Expression Language (${...}) and JSTL tags (<c:forEach>). Containerized Docker deployment.", C_NAVY)
    ]

    uw = Inches(11.733)
    uh = Inches(0.88)
    for i, (utitle, udesc, ucol) in enumerate(units):
        uy = Inches(1.8) + i * (uh + Inches(0.18))
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), uy, uw, uh)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        # Left color badge
        badge = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), uy + Inches(0.14), Inches(2.6), Inches(0.6))
        badge.fill.solid()
        badge.fill.fore_color.rgb = ucol
        badge.line.fill.background()
        btf = badge.text_frame
        btf.margin_top = Inches(0.12)
        bp = btf.paragraphs[0]
        bp.text = utitle
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = C_WHITE
        bp.alignment = PP_ALIGN.CENTER

        # Description text
        dtb = s5.shapes.add_textbox(Inches(3.75), uy + Inches(0.12), Inches(8.6), Inches(0.65))
        dtf = dtb.text_frame
        dtf.word_wrap = True
        dp = dtf.paragraphs[0]
        dp.text = udesc
        dp.font.size = Pt(12)
        dp.font.color.rgb = C_TEXT_DARK

    # =========================================================================
    # SLIDE 6: OOP Design & Polymorphic Pricing Engine
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Object-Oriented Programming (Unit I)", "Polymorphism, Inheritance & Abstraction in Pricing")

    oop_cards = [
        ("Abstract Base: ServicePlan", C_BLUE, [
            "Encapsulates common plan attributes: name, setupFee, monthlyRetainer, billingCycle.",
            "Defines abstract method: public abstract PricingQuote calculateTotal(boolean annual, ...)",
            "Guarantees a unified contract across all current and future service plan variants."
        ]),
        ("Concrete Subclasses", C_TEAL, [
            "StarterPlan: Specialized for single-page MSME launch (₹5,000 setup + ₹1,500/mo retainer).",
            "GrowthPlan: 5-page dynamic CMS website with SEO (₹12,000 setup + ₹3,000/mo retainer).",
            "EnterprisePlan: Custom multi-page web app with priority SLA (₹25,000 setup + ₹6,000/mo)."
        ]),
        ("Dynamic Polymorphic Dispatch", C_EMERALD, [
            "PricingService dynamically instantiates the correct plan based on client inputs.",
            "At runtime, Java executes the overridden calculateTotal() method dynamically.",
            "Applies 15% annual billing discount and optional add-on configurations cleanly."
        ])
    ]

    for i, (title, col, items) in enumerate(oop_cards):
        ox = Inches(0.8) + i * (cw + Inches(0.4))
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ox, Inches(1.8), cw, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        ib = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, ox, Inches(1.8), cw, Inches(0.12))
        ib.fill.solid()
        ib.fill.fore_color.rgb = col
        ib.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.35)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(14)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_MUTED
            p.space_after = Pt(10)

    # =========================================================================
    # SLIDE 7: Multithreading & Proactive Monitoring Engine
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Multithreading & Networking (Unit II)", "Proactive Fleet Telemetry & Background Daemon")

    mt_items = [
        ("Runnable Daemon Architecture", "MonitorService implements Runnable and HealthChecker. It runs as a setDaemon(true) background worker thread initialized during Tomcat boot via AppContextListener.", C_BLUE),
        ("Autonomous Health Sweeps", "Every 45 seconds, the daemon sweeps all client domains in the SQLite database, issuing HttpURLConnection HTTP HEAD/GET probes and measuring round-trip latency.", C_TEAL),
        ("Automated Issue Resolution", "If latency spikes above 1200ms or response codes indicate errors, the service automatically logs an Urgent Maintenance Ticket in the SQLite database without user intervention.", C_AMBER),
        ("Omnichannel Incident Alerts", "Integrates Jakarta Mail (SMTP) to automatically dispatch incident alert emails to the business owner, notifying them before their customers experience disruption.", C_EMERALD)
    ]

    for i, (title, desc, col) in enumerate(mt_items):
        mx = Inches(0.8) + (i % 2) * (fw + Inches(0.4))
        my = Inches(1.8) + (i // 2) * (fh + Inches(0.35))

        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mx, my, fw, fh)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.25)
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 8: Database Architecture & Concurrency
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Database & Persistence (Unit III)", "SQLite Relational Schema & Multi-Threaded Concurrency")

    db_cards = [
        ("Core Relational Tables", C_NAVY, [
            "customers: Verified business accounts created via Firebase Google OAuth.",
            "website_orders: Generated project bookings with selected tier & pricing.",
            "sites: Monitored client domains with live uptime, latency, & speed metrics.",
            "tickets: Support tickets logged by clients or auto-generated by the monitor.",
            "leads: Instant website audits and inbound business inquiries."
        ]),
        ("Concurrency Engineering", C_BLUE, [
            "WAL Journal Mode: PRAGMA journal_mode = WAL enables simultaneous readers and writers.",
            "Busy Timeout: PRAGMA busy_timeout = 10000 eliminates SQLite lock collisions.",
            "DAO Pattern: Each query allocates its own connection via try-with-resources.",
            "Thread Isolation: Tested with 120 simultaneous queries across 10 threads — 0 failures."
        ]),
        ("Data Access Object (DAO)", C_TEAL, [
            "SiteDAO: Full CRUD for monitored fleet telemetry.",
            "CustomerDAO: Customer profiles & Google OAuth records.",
            "OrderDAO: Order creation and invoice status management.",
            "TicketDAO: Ticket creation, lifecycle filtering, and resolution.",
            "LeadDAO: Free audit ingestion and lead nurturing."
        ])
    ]

    for i, (title, col, items) in enumerate(db_cards):
        dx = Inches(0.8) + i * (cw + Inches(0.4))
        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, dx, Inches(1.8), cw, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        ib = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, dx, Inches(1.8), cw, Inches(0.12))
        ib.fill.solid()
        ib.fill.fore_color.rgb = col
        ib.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.35)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(14)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_MUTED
            p.space_after = Pt(10)

    # =========================================================================
    # SLIDE 9: User Workflow & Key Modules
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_header(s9, "Application Workflow", "End-to-End Customer & Operations Journey")

    steps = [
        ("Step 1: Free Audit", "Owner enters domain URL to evaluate security, SSL, mobile readiness, and speed score.", C_BLUE),
        ("Step 2: Transparent Pricing", "Interactive slider calculates customized package quotes for Starter, Growth, or Enterprise tiers.", C_TEAL),
        ("Step 3: Google Sign-In", "1-click Google authentication via Firebase popup. Eliminates password fatigue and friction.", C_EMERALD),
        ("Step 4: Expert Consultation", "Post-order consultation page connects the customer directly via WhatsApp or Priority Direct Call.", C_AMBER),
        ("Step 5: Portal & Telemetry", "Client tracks ongoing site health, uptime percentages, and files maintenance requests in real-time.", C_NAVY)
    ]

    sw = Inches(2.18)
    for i, (title, desc, col) in enumerate(steps):
        sx = Inches(0.8) + i * (sw + Inches(0.2))
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, Inches(1.8), sw, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        # Step circle
        circ = s9.shapes.add_shape(MSO_SHAPE.OVAL, sx + Inches(0.74), Inches(2.1), Inches(0.7), Inches(0.7))
        circ.fill.solid()
        circ.fill.fore_color.rgb = col
        circ.line.fill.background()
        ctf = circ.text_frame
        ctf.margin_top = Inches(0.12)
        cp = ctf.paragraphs[0]
        cp.text = str(i + 1)
        cp.font.size = Pt(16)
        cp.font.bold = True
        cp.font.color.rgb = C_WHITE
        cp.alignment = PP_ALIGN.CENTER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(1.2)
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(10)
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 10: Security, DevOps & Cloud Deployment
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10)
    add_header(s10, "Security & DevOps", "Cloud Containerization, Authentication & Deployment")

    sec_cards = [
        ("Firebase Google OAuth 2.0", C_BLUE, [
            "Client executes popup sign-in with prompt: 'select_account'.",
            "Firebase client passes cryptographic JWT ID token to Tomcat backend.",
            "Backend parses token payload to verify email, name, and identity securely.",
            "Zero plaintext passwords stored in the database."
        ]),
        ("Docker Containerization", C_TEAL, [
            "Base image: tomcat:9.0-jdk17-temurin for modern Java 17 execution.",
            "Web content copied into ROOT so the webapp serves directly at root path (/).",
            "Dynamic port binding: script binds Tomcat to Render's $PORT (10000).",
            "Shutdown port disabled (port=-1) to prevent probe conflicts."
        ]),
        ("Production Cloud on Render", C_EMERALD, [
            "Automated CI/CD linked to GitHub repository (TakeItEasySystem/java-MiniProject).",
            "Automatic container builds and health check verification.",
            "Free SSL HTTPS certificate issued automatically for web encryption.",
            "Memory constrained to 320MB JVM heap to ensure stability on Free Tier."
        ])
    ]

    for i, (title, col, items) in enumerate(sec_cards):
        sx = Inches(0.8) + i * (cw + Inches(0.4))
        card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, Inches(1.8), cw, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        ib = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, sx, Inches(1.8), cw, Inches(0.12))
        ib.fill.solid()
        ib.fill.fore_color.rgb = col
        ib.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.35)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(14)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_MUTED
            p.space_after = Pt(10)

    # =========================================================================
    # SLIDE 11: Testing, Verification & Key Metrics
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_bg(s11)
    add_header(s11, "Validation & Quality Assurance", "Automated Testing, Concurrency Verification & Metrics")

    metrics = [
        ("100%", "Test Suite Pass", "Automated BackendTest.java verifies DB schema, DAOs, Pricing, and Telemetry.", C_EMERALD),
        ("120", "Concurrent Queries", "Stress tested across 10 simultaneous threads in SQLite WAL mode with 0 errors.", C_BLUE),
        ("< 300ms", "Avg API Response", "Lightweight Servlets and optimized queries deliver lightning-fast page loads.", C_TEAL),
        ("Zero", "Password Breaches", "100% credential delegation to Google OAuth guarantees zero credential leaks.", C_AMBER)
    ]

    mw = Inches(2.73)
    for i, (stat, label, detail, col) in enumerate(metrics):
        mx = Inches(0.8) + i * (mw + Inches(0.26))
        card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mx, Inches(1.8), mw, Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.25)
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = stat
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = col
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_DARK
        p2.alignment = PP_ALIGN.CENTER
        p2.space_after = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = detail
        p3.font.size = Pt(10)
        p3.font.color.rgb = C_TEXT_MUTED
        p3.alignment = PP_ALIGN.CENTER

    # Bottom Summary Box
    bot = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.35), Inches(11.733), Inches(2.55))
    bot.fill.solid()
    bot.fill.fore_color.rgb = C_WHITE
    bot.line.color.rgb = C_BORDER
    bot.line.width = Pt(1.5)

    btf = bot.text_frame
    btf.word_wrap = True
    btf.margin_top = Inches(0.3)
    btf.margin_left = Inches(0.35)
    btf.margin_right = Inches(0.35)

    bp = btf.paragraphs[0]
    bp.text = "Key Technical Milestones Achieved"
    bp.font.size = Pt(16)
    bp.font.bold = True
    bp.font.color.rgb = C_TEXT_DARK
    bp.space_after = Pt(10)

    points = [
        "Eliminated SQLite Database Lockups: Transitioned from singleton connection to thread-safe connection-per-request model with WAL mode.",
        "Engineered Seamless Cloud Port Mapping: Enabled Tomcat to dynamically bind to Render's required PORT environment variable.",
        "Implemented Enterprise Post-Order Flow: Direct routing from order placement to consultation channel selection (WhatsApp / Direct Call).",
        "Full Syllabus Compliance: Satisfied all Full Stack Java Programming (FSJP) competencies from Core OOP to Cloud Deployment."
    ]
    for pt in points:
        bp2 = btf.add_paragraph()
        bp2.text = "✓ " + pt
        bp2.font.size = Pt(12)
        bp2.font.color.rgb = C_TEXT_MUTED
        bp2.space_after = Pt(6)

    # =========================================================================
    # SLIDE 12: Future Roadmap & Conclusion (Dark Theme)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_bg(s12, dark=True)

    # Header
    tb12 = s12.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
    tf12 = tb12.text_frame
    p = tf12.paragraphs[0]
    p.text = "FUTURE ENHANCEMENTS & CONCLUSION"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL
    p.space_after = Pt(4)
    p2 = tf12.add_paragraph()
    p2.text = "Scalability Roadmap & Concluding Remarks"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE

    # 3 Future Scope Cards
    fscope = [
        ("AI Autonomous Self-Healing", [
            "Auto-purge CDN cache upon latency spikes.",
            "Automated SSL renewal triggers via Let's Encrypt API.",
            "Smart AI diagnosis for broken form inputs."
        ]),
        ("FinTech Payment Integration", [
            "Razorpay & Stripe recurring subscription webhooks.",
            "Automated GST invoice generation (PDF).",
            "Self-service billing updates on Client Portal."
        ]),
        ("Multi-Region Edge Nodes", [
            "Distributed ping collectors across Mumbai, Singapore, and Frankfurt.",
            "Real-time geographic latency map in Admin Console.",
            "High-availability MySQL / PostgreSQL cluster."
        ])
    ]

    for i, (title, items) in enumerate(fscope):
        fx = Inches(0.8) + i * (cw + Inches(0.4))
        card = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx, Inches(1.8), cw, Inches(3.0))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_DARK
        card.line.color.rgb = C_BORDER_DARK
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.25)
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_TEAL
        p.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(11)
            p.font.color.rgb = RGBColor(203, 213, 225)
            p.space_after = Pt(6)

    # Thank you footer card
    tcard = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.1), Inches(11.733), Inches(1.8))
    tcard.fill.solid()
    tcard.fill.fore_color.rgb = C_BLUE
    tcard.line.fill.background()

    tt = tcard.text_frame
    tt.word_wrap = True
    tt.margin_top = Inches(0.25)
    tt.margin_left = Inches(0.4)
    tt.margin_right = Inches(0.4)

    p = tt.paragraphs[0]
    p.text = "Thank You! Questions & Answers Welcome."
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(6)

    p2 = tt.add_paragraph()
    p2.text = "Presented by: Ayush Jha (A18)  •  Rishikant Mishra (A29)  •  Yadnesh Sawant (A68)  |  Thakur Shree DPS College"
    p2.font.size = Pt(13)
    p2.font.color.rgb = RGBColor(224, 231, 255)

    p3 = tt.add_paragraph()
    p3.text = "Live Repository: https://github.com/TakeItEasySystem/java-MiniProject"
    p3.font.size = Pt(11)
    p3.font.color.rgb = RGBColor(199, 210, 254)

    output_path = r"C:\java mini project\WebEase_Project_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    create_deck()
